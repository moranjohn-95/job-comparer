import json
import re
from collections.abc import Callable, Iterator
from io import BytesIO
from uuid import uuid4

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from comparison import MAX_CV_CHARS, source_excerpts
from cv_upload import MAX_UPLOAD_BYTES
from database import get_engine
from main import MAX_CV_LENGTH, app
from models import SavedCV, User

PASSWORD = "correct-horse-battery-123"


def make_pdf(*, text: str | None = None, image: bool = False) -> bytes:
    if image:
        resources = b"/XObject << /Im1 4 0 R >>"
        image_data = b"\xff\x00\x00"
        fourth = (
            b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 "
            b"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\n"
            b"stream\n" + image_data + b"\nendstream"
        )
        content = b"q 100 0 0 100 72 600 cm /Im1 Do Q"
    else:
        resources = b"/Font << /F1 4 0 R >>"
        fourth = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
        content = (
            b"BT /F1 12 Tf 72 720 Td ("
            + (text or "").encode("ascii")
            + b") Tj ET"
        )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << " + resources + b" >> /Contents 5 0 R >>",
        fourth,
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    startxref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        b"trailer\n<< /Size "
        + str(len(offsets)).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(startxref).encode()
        + b"\n%%EOF\n"
    )
    return bytes(output)


def make_docx(text: str) -> bytes:
    document = Document()
    if text:
        document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def upload(
    client: TestClient, headers: dict[str, str], filename: str, data: bytes
):
    return client.post(
        "/cv/upload", files={"file": (filename, data)}, headers=headers
    )


@pytest.fixture
def client() -> Iterator[TestClient]:
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def create_user(
    client: TestClient,
) -> Iterator[Callable[[], tuple[int, dict[str, str]]]]:
    emails: list[str] = []

    def create() -> tuple[int, dict[str, str]]:
        email = f"cv-test-{uuid4().hex}@example.com"
        emails.append(email)
        signup = client.post(
            "/signup", json={"email": email, "password": PASSWORD}
        )
        assert signup.status_code == 201
        login = client.post(
            "/login", json={"email": email, "password": PASSWORD}
        )
        assert login.status_code == 200
        return signup.json()["id"], {
            "Authorization": f"Bearer {login.json()['access_token']}"
        }

    yield create
    with Session(get_engine()) as session:
        session.execute(delete(User).where(User.email.in_(emails)))
        session.commit()


def test_cv_persists_across_requests_and_in_database(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    user_id, headers = create_user()
    cv_text = "Software engineer\nPython and PostgreSQL experience"

    saved = client.put("/cv", json={"text": cv_text}, headers=headers)
    retrieved = client.get("/cv", headers=headers)

    assert saved.status_code == 200
    assert saved.json() == {"text": cv_text}
    assert retrieved.status_code == 200
    assert retrieved.json() == {"text": cv_text}
    with Session(get_engine()) as session:
        assert session.get(SavedCV, user_id).text == cv_text


def test_saving_again_replaces_the_only_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    user_id, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "First version"}, headers=headers
        ).status_code
        == 200
    )

    replaced = client.put(
        "/cv", json={"text": "Updated version"}, headers=headers
    )

    assert replaced.status_code == 200
    assert client.get("/cv", headers=headers).json() == {
        "text": "Updated version"
    }
    with Session(get_engine()) as session:
        count = session.scalar(
            select(func.count())
            .select_from(SavedCV)
            .where(SavedCV.user_id == user_id)
        )
        assert count == 1


def test_delete_removes_the_saved_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    user_id, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "To delete"}, headers=headers
        ).status_code
        == 200
    )

    deleted = client.delete("/cv", headers=headers)
    missing = client.get("/cv", headers=headers)

    assert deleted.status_code == 204
    assert missing.status_code == 404
    assert missing.json() == {"detail": "CV not found"}
    with Session(get_engine()) as session:
        assert session.get(SavedCV, user_id) is None


@pytest.mark.parametrize("method", ["GET", "PUT", "DELETE"])
def test_cv_endpoints_require_authentication(
    client: TestClient, method: str
) -> None:
    body = {"text": "A valid CV"} if method == "PUT" else None

    response = client.request(method, "/cv", json=body)

    assert response.status_code == 401
    assert response.json() == {"detail": "Authentication required"}


def test_users_cannot_read_or_change_each_others_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    first_id, first_headers = create_user()
    _, second_headers = create_user()
    assert (
        client.put(
            "/cv",
            json={"text": "First user's private CV"},
            headers=first_headers,
        ).status_code
        == 200
    )

    assert client.get("/cv", headers=second_headers).status_code == 404
    spoofed = client.put(
        "/cv",
        json={"text": "Attempted overwrite", "user_id": first_id},
        headers=second_headers,
    )
    assert spoofed.status_code == 422
    assert client.get("/cv", headers=second_headers).status_code == 404

    assert (
        client.put(
            "/cv", json={"text": "Second user's CV"}, headers=second_headers
        ).status_code
        == 200
    )
    assert client.get("/cv", headers=first_headers).json() == {
        "text": "First user's private CV"
    }
    assert client.get("/cv", headers=second_headers).json() == {
        "text": "Second user's CV"
    }
    assert client.delete("/cv", headers=second_headers).status_code == 204
    assert client.get("/cv", headers=first_headers).json() == {
        "text": "First user's private CV"
    }


@pytest.mark.parametrize(
    "invalid_text,detail",
    [
        ("", "CV text must not be blank"),
        (" \t\n ", "CV text must not be blank"),
        (
            "x" * (MAX_CV_LENGTH + 1),
            f"CV text exceeds {MAX_CV_LENGTH} characters",
        ),
    ],
    ids=["empty", "whitespace", "too-long"],
)
def test_invalid_cv_does_not_replace_existing_text(
    client: TestClient,
    create_user: Callable[[], tuple[int, dict[str, str]]],
    invalid_text: str,
    detail: str,
) -> None:
    _, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "Valid CV"}, headers=headers
        ).status_code
        == 200
    )

    response = client.put("/cv", json={"text": invalid_text}, headers=headers)

    assert response.status_code == 422
    assert response.json() == {"detail": detail}
    assert client.get("/cv", headers=headers).json() == {"text": "Valid CV"}


def test_pdf_upload_extracts_text_and_replaces_saved_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    user_id, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "Original CV"}, headers=headers
        ).status_code
        == 200
    )

    response = upload(
        client, headers, "resume.PDF", make_pdf(text="PDF resume text")
    )

    assert response.status_code == 200
    assert response.json() == {"text": "PDF resume text"}
    assert client.get("/cv", headers=headers).json() == {
        "text": "PDF resume text"
    }
    with Session(get_engine()) as session:
        assert session.get(SavedCV, user_id).text == "PDF resume text"
        assert (
            session.scalar(
                select(func.count())
                .select_from(SavedCV)
                .where(SavedCV.user_id == user_id)
            )
            == 1
        )


def test_docx_upload_extracts_paragraphs_and_tables(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    _, headers = create_user()
    document = Document()
    document.add_paragraph("DOCX resume text")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "PostgreSQL"
    output = BytesIO()
    document.save(output)

    response = upload(client, headers, "resume.docx", output.getvalue())

    assert response.status_code == 200
    assert response.json() == {"text": "DOCX resume text\nPython\tPostgreSQL"}
    assert client.get("/cv", headers=headers).json() == response.json()


@pytest.mark.parametrize("source", ["text", "pdf", "docx"])
def test_all_cv_sections_reach_provider_without_truncation(
    client: TestClient,
    create_user: Callable[[], tuple[int, dict[str, str]]],
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
    source: str,
) -> None:
    user_id, headers = create_user()
    sections = [
        "Work experience: Built internal reporting tools.",
        "Education: Diploma in software development.",
        "Technical skills: Python, SQL, scikit-learn.",
        "Training: Completed a practical statistics course.",
        "Projects: Trained a machine-learning random forest classifier.",
    ]
    cv_text = "\n".join(sections)
    if source == "text":
        # Keep meaningful evidence at the very end of the accepted input.
        cv_text = " " * (MAX_CV_CHARS - len(cv_text)) + cv_text
        saved = client.put("/cv", json={"text": cv_text}, headers=headers)
    else:
        output = BytesIO()
        if source == "pdf":
            writer = PdfWriter()
            for section in sections:
                page = PdfReader(BytesIO(make_pdf(text=section))).pages[0]
                writer.add_page(page)
            writer.write(output)
        else:
            document = Document()
            for section in sections[:-1]:
                document.add_paragraph(section)
            # Project evidence in the final table must also survive.
            document.add_table(rows=1, cols=1).cell(0, 0).text = sections[-1]
            document.save(output)
        saved = upload(
            client, headers, f"synthetic.{source}", output.getvalue(),
        )

    assert saved.status_code == 200
    assert saved.json()["text"] == cv_text
    assert client.get("/cv", headers=headers).json()["text"] == cv_text
    with Session(get_engine()) as session:
        assert session.get(SavedCV, user_id).text == cv_text

    job_text = "Machine learning AND deep learning experience required."
    job = client.post(
        "/jobs", headers=headers,
        json={
            "title": "Engineer", "company_name": "Synthetic Company",
            "description": job_text,
        },
    )
    assert job.status_code == 201
    monkeypatch.setenv("AI_COMPARISON_ENABLED", "true")
    monkeypatch.setenv("OPENAI_API_KEY", "test-provider-key")
    monkeypatch.setenv("OPENAI_MODEL", "gpt-6.1-sol")
    payloads = []

    def fake_post(*_: object, **kwargs: object) -> httpx.Response:
        payloads.append(kwargs["json"])
        return httpx.Response(
            200,
            request=httpx.Request("POST", "https://example.test"),
            json={
                "status": "completed",
                "output": [{
                    "type": "message",
                    "content": [{
                        "type": "output_text",
                        "text": json.dumps({
                            "inventory_complete": True,
                            "requirements": [], "assessments": [],
                        }),
                    }],
                }],
            },
        )

    monkeypatch.setattr("comparison._post_once", fake_post)
    response = client.post(
        f"/jobs/{job.json()['id']}/compare", headers=headers,
    )

    assert response.status_code == 200
    assert len(payloads) == 1
    assert json.loads(payloads[0]["input"][1]["content"]) == {
        "cv_excerpts": [
            {**item, "text": re.sub(r"\s+", " ", item["text"])}
            for item in source_excerpts(cv_text, "cv")
        ],
        "job_excerpts": source_excerpts(job_text, "job"),
    }
    for section in sections:
        assert section not in caplog.text


@pytest.mark.parametrize(
    "filename,data,status_code,detail",
    [
        (
            "resume.txt",
            b"Plain text CV",
            415,
            "Unsupported file type; upload a PDF or DOCX",
        ),
        (
            "resume.pdf",
            b"This is not a PDF",
            415,
            "File content does not match PDF format",
        ),
        (
            "resume.docx",
            make_pdf(text="Wrong format"),
            415,
            "File content does not match DOCX format",
        ),
        ("resume.pdf", b"%PDF-1.4\ninvalid", 422, "Damaged PDF file"),
        ("resume.docx", b"PK\x03\x04invalid", 422, "Damaged DOCX file"),
        (
            "resume.pdf",
            make_pdf(image=True),
            422,
            "Scanned PDF has no selectable text; OCR is not supported",
        ),
        ("resume.pdf", make_pdf(), 422, "PDF contains no readable text"),
        ("resume.docx", make_docx(""), 422, "DOCX contains no readable text"),
        ("resume.pdf", b"", 422, "Uploaded file is empty"),
    ],
    ids=[
        "unsupported-extension",
        "pdf-content-mismatch",
        "docx-content-mismatch",
        "damaged-pdf",
        "damaged-docx",
        "scanned-pdf",
        "text-empty-pdf",
        "text-empty-docx",
        "empty-file",
    ],
)
def test_invalid_upload_preserves_saved_cv(
    client: TestClient,
    create_user: Callable[[], tuple[int, dict[str, str]]],
    filename: str,
    data: bytes,
    status_code: int,
    detail: str,
) -> None:
    _, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "Keep this CV"}, headers=headers
        ).status_code
        == 200
    )

    response = upload(client, headers, filename, data)

    assert response.status_code == status_code
    assert response.json() == {"detail": detail}
    assert client.get("/cv", headers=headers).json() == {
        "text": "Keep this CV"
    }


def test_password_protected_pdf_is_rejected_without_replacing_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    _, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "Keep this CV"}, headers=headers
        ).status_code
        == 200
    )
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)

    response = upload(client, headers, "locked.pdf", output.getvalue())

    assert response.status_code == 422
    assert response.json() == {
        "detail": "Password-protected PDF files are not supported"
    }
    assert client.get("/cv", headers=headers).json() == {
        "text": "Keep this CV"
    }


def test_upload_size_limit_preserves_saved_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    _, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "Keep this CV"}, headers=headers
        ).status_code
        == 200
    )

    response = upload(
        client, headers, "large.pdf", b"%PDF-" + b"x" * MAX_UPLOAD_BYTES
    )

    assert response.status_code == 413
    assert response.json() == {"detail": "File exceeds the 5 MiB upload limit"}
    assert client.get("/cv", headers=headers).json() == {
        "text": "Keep this CV"
    }


def test_extracted_text_limit_preserves_saved_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    _, headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "Keep this CV"}, headers=headers
        ).status_code
        == 200
    )

    response = upload(
        client, headers, "long.docx", make_docx("x" * (MAX_CV_LENGTH + 1))
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": f"CV text exceeds {MAX_CV_LENGTH} characters"
    }
    assert client.get("/cv", headers=headers).json() == {
        "text": "Keep this CV"
    }


def test_upload_requires_authentication(client: TestClient) -> None:
    response = upload(client, {}, "resume.pdf", make_pdf(text="Private CV"))

    assert response.status_code == 401
    assert response.json() == {"detail": "Authentication required"}


def test_upload_replaces_only_the_authenticated_users_cv(
    client: TestClient, create_user: Callable[[], tuple[int, dict[str, str]]]
) -> None:
    _, first_headers = create_user()
    _, second_headers = create_user()
    assert (
        client.put(
            "/cv", json={"text": "First user's CV"}, headers=first_headers
        ).status_code
        == 200
    )
    assert client.get("/cv", headers=second_headers).status_code == 404

    response = upload(
        client,
        second_headers,
        "second.docx",
        make_docx("Second user's uploaded CV"),
    )

    assert response.status_code == 200
    assert client.get("/cv", headers=second_headers).json() == {
        "text": "Second user's uploaded CV"
    }
    assert client.get("/cv", headers=first_headers).json() == {
        "text": "First user's CV"
    }
